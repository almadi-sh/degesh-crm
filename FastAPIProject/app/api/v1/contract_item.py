import logging
from io import BytesIO
from typing import List
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy.orm import Session

from app.api.deps import get_db
from app.models.contract import Contract
from app.models.contract_item import ContractItem
from app.models.inventory import Inventory
from app.models.product import Product
from app.models.reservation import Reservation
from app.schemas.contract_item import ContractItemCreate, ContractItemOut, ContractItemUpdate

router = APIRouter(prefix="/contract-items", tags=["Contract Items"])
logger = logging.getLogger("uvicorn.access")


def _calculate_total_amount(quantity: float, price: float, vat_enabled: bool) -> float:
    vat_multiplier = 1.16 if vat_enabled else 1.0
    return quantity * price * vat_multiplier


def _get_inventory(db: Session, product_id: int) -> Inventory:
    inv = db.query(Inventory).filter(Inventory.product_id == product_id).first()
    if not inv:
        raise HTTPException(status_code=404, detail=f"Inventory not found for product_id {product_id}")
    return inv


def _reserve_inventory_delta(db: Session, product_id: int, quantity_delta: float) -> Inventory:
    inv = _get_inventory(db, product_id)
    if quantity_delta > 0 and inv.quantity_available < quantity_delta:
        raise HTTPException(status_code=400, detail=f"Not enough inventory for product_id {product_id}")
    inv.quantity_available -= quantity_delta
    inv.quantity_reserved += quantity_delta
    if inv.quantity_available < 0:
        raise HTTPException(status_code=400, detail=f"Inventory cannot be negative for product_id {product_id}")
    if inv.quantity_reserved < 0:
        inv.quantity_reserved = 0
    db.add(inv)
    return inv


def _upsert_reservation(
    db: Session,
    contract: Contract,
    contract_item: ContractItem,
    quantity: float,
    product_id: int,
    status: str = "active",
) -> Reservation:
    reservation = db.query(Reservation).filter(Reservation.contract_item_id == contract_item.id).first()
    inventory = _get_inventory(db, product_id)
    if reservation:
        reservation.inventory_id = inventory.id
        reservation.product_id = product_id
        reservation.contract_id = contract.id
        reservation.contract_item_id = contract_item.id
        reservation.employee_id = contract.owner_employee_id
        reservation.customer_id = contract.customer_id
        reservation.quantity = quantity
        reservation.priority = contract_item.appendix_number
        reservation.delivery_due_date = None
        reservation.status = status
    else:
        reservation = Reservation(
            inventory_id=inventory.id,
            product_id=product_id,
            contract_id=contract.id,
            contract_item_id=contract_item.id,
            employee_id=contract.owner_employee_id,
            customer_id=contract.customer_id,
            quantity=quantity,
            priority=contract_item.appendix_number,
            delivery_due_date=None,
            status=status,
        )
    db.add(reservation)
    return reservation


def _build_appendix_xlsx(
    rows: list[list[str | float]],
    header_title: str,
    parties_line: str,
    seller_signature: list[str],
    buyer_signature: list[str],
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Appendix"

    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = header_title
    worksheet["A1"].font = Font(bold=True)
    worksheet["A1"].alignment = Alignment(horizontal="left")

    worksheet.merge_cells("A2:F2")
    worksheet["A2"] = parties_line
    worksheet["A2"].font = Font(bold=True)
    worksheet["A2"].alignment = Alignment(horizontal="left")

    worksheet.merge_cells("A4:F4")
    worksheet["A4"] = "Спецификация"
    worksheet["A4"].font = Font(bold=True)
    worksheet["A4"].alignment = Alignment(horizontal="center")

    thin_side = Side(style="thin")
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    table_start_row = 5
    for index, row in enumerate(rows):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=table_start_row + index, column=column_index, value=value)
            cell.border = table_border
            if index == 0:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
            else:
                horizontal_alignment = "left" if column_index == 1 else "center"
                cell.alignment = Alignment(horizontal=horizontal_alignment, vertical="center", wrap_text=True)

    footer_start_row = table_start_row + len(rows) + 2
    worksheet.cell(row=footer_start_row, column=1, value="Продавец:")
    worksheet.cell(row=footer_start_row + 1, column=1, value=seller_signature[0])
    worksheet.cell(row=footer_start_row + 2, column=1, value=seller_signature[1])

    buyer_column = 6
    worksheet.cell(row=footer_start_row, column=buyer_column, value="Покупатель:")
    worksheet.cell(row=footer_start_row + 1, column=buyer_column, value=buyer_signature[0])
    worksheet.cell(row=footer_start_row + 2, column=buyer_column, value=buyer_signature[1])

    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 24
    worksheet.column_dimensions["D"].width = 24
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.post("/", response_model=ContractItemOut)
def create_contract_item(data: ContractItemCreate, db: Session = Depends(get_db)):
    contract = db.query(Contract).filter(Contract.id == data.contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found")

    if contract.status != "Draft":
        raise HTTPException(status_code=400, detail="Contract items can only be added in Draft status")

    _reserve_inventory_delta(db, data.product_id, data.quantity)
    total_amount = _calculate_total_amount(data.quantity, data.price, data.vat_enabled)

    contract_item = ContractItem(
        contract_id=data.contract_id,
        product_id=data.product_id,
        quantity=data.quantity,
        price=data.price,
        total_amount=total_amount,
        vat_enabled=data.vat_enabled,
        delivery_enabled=data.delivery_enabled,
        delivery_terms=data.delivery_terms if data.delivery_enabled else None,
        appendix_number=max(data.appendix_number, 1),
    )
    db.add(contract_item)
    db.flush()
    _upsert_reservation(db, contract, contract_item, data.quantity, data.product_id, status="active")
    db.commit()
    db.refresh(contract_item)
    logger.info("POST /contract-items -> %s", contract_item.id)
    return contract_item


@router.get("/", response_model=List[ContractItemOut])
def list_contract_items(
    db: Session = Depends(get_db),
    customer_id: int | None = Query(default=None),
    contract_id: int | None = Query(default=None),
):
    query = db.query(ContractItem)
    if contract_id is not None:
        query = query.filter(ContractItem.contract_id == contract_id)
    if customer_id is not None:
        query = query.join(Contract).filter(Contract.customer_id == customer_id)
    return query.all()


@router.get("/contract/{contract_id}/appendix/{appendix_number}/export-xlsx")
def export_appendix_xlsx(contract_id: int, appendix_number: int, db: Session = Depends(get_db)):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found")

    items = (
        db.query(ContractItem, Product)
        .join(Product, Product.id == ContractItem.product_id)
        .filter(ContractItem.contract_id == contract_id, ContractItem.appendix_number == appendix_number)
        .all()
    )
    if not items:
        raise HTTPException(status_code=404, detail="Appendix has no items")

    rows: list[list[str | float]] = [
        [
            "Продукт",
            "Количество, л/кг/п.е.",
            "Цена вкл. НДС и таможенную пошлину, тенге/л/кг",
            "Общая стоимость, тенге",
            "Срок поставки",
            "Срок оплаты",
        ]
    ]
    for item, product in items:
        price_with_vat = item.price * 1.16 if item.vat_enabled else item.price
        rows.append(
            [
                product.name,
                item.quantity,
                price_with_vat,
                item.total_amount,
                item.delivery_terms or "",
                "",
            ]
        )

    contract_date = contract.contract_date.strftime("%d.%m.%Y")
    customer_name = contract.customer.name if contract.customer else "Покупатель"
    signatures = contract.contract_document.get("signatures", {}) if contract.contract_document else {}
    seller_name = signatures.get("linked_customer_name") or "ТОО DeGesH"
    seller_role = signatures.get("seller_position") or "Директор"
    seller_signer_name = signatures.get("seller_name") or ""
    buyer_role = signatures.get("buyer_position") or (contract.customer.contract_signer_role if contract.customer else "") or "Директор"
    buyer_signer_name = signatures.get("buyer_name") or (contract.customer.contract_signer_full_name if contract.customer else "") or ""

    file_bytes = _build_appendix_xlsx(
        rows=rows,
        header_title=f"Приложение {appendix_number} к договору {contract.contract_number} от {contract_date} года",
        parties_line=f"между {seller_name} и {customer_name}",
        seller_signature=[f"{seller_role} {seller_name}", seller_signer_name],
        buyer_signature=[f"{buyer_role} {customer_name}", buyer_signer_name],
    )
    filename = f"appendix-{appendix_number}-contract-{contract.contract_number}.xlsx"
    ascii_filename = f"appendix-{appendix_number}-contract.xlsx"
    content_disposition = (
        f"attachment; filename={ascii_filename}; filename*=UTF-8''{quote(filename)}"
    )
    return Response(
        content=file_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition},
    )


@router.put("/{contract_item_id}", response_model=ContractItemOut)
def update_contract_item(
    contract_item_id: int,
    data: ContractItemUpdate,
    db: Session = Depends(get_db),
):
    contract_item = db.query(ContractItem).filter(ContractItem.id == contract_item_id).first()
    if not contract_item:
        raise HTTPException(status_code=404, detail="Contract item not found")

    if contract_item.contract.status != "Draft":
        raise HTTPException(status_code=400, detail="Contract items can only be updated in Draft status")

    old_product_id = contract_item.product_id
    old_quantity = contract_item.quantity
    update_data = data.dict(exclude_unset=True)
    if update_data.get("delivery_enabled") is False:
        update_data["delivery_terms"] = None

    new_product_id = update_data.get("product_id", old_product_id)
    new_quantity = update_data.get("quantity", old_quantity)

    if new_product_id != old_product_id:
        _reserve_inventory_delta(db, new_product_id, new_quantity)
        _reserve_inventory_delta(db, old_product_id, -old_quantity)
    elif new_quantity != old_quantity:
        _reserve_inventory_delta(db, old_product_id, new_quantity - old_quantity)

    for key, value in update_data.items():
        setattr(contract_item, key, value)

    if "quantity" in update_data or "price" in update_data or "vat_enabled" in update_data:
        contract_item.total_amount = _calculate_total_amount(
            contract_item.quantity,
            contract_item.price,
            contract_item.vat_enabled,
        )

    _upsert_reservation(
        db,
        contract_item.contract,
        contract_item,
        contract_item.quantity,
        contract_item.product_id,
        status="active",
    )
    db.commit()
    db.refresh(contract_item)
    logger.info("PUT /contract-items/%s", contract_item_id)
    return contract_item


@router.delete("/{contract_item_id}", status_code=204)
def delete_contract_item(contract_item_id: int, db: Session = Depends(get_db)):
    contract_item = db.query(ContractItem).filter(ContractItem.id == contract_item_id).first()
    if not contract_item:
        raise HTTPException(status_code=404, detail="Contract item not found")

    _reserve_inventory_delta(db, contract_item.product_id, -contract_item.quantity)

    db.query(Reservation).filter(Reservation.contract_item_id == contract_item.id).delete(
        synchronize_session=False,
    )

    db.delete(contract_item)
    db.commit()
    logger.info("DELETE /contract-items/%s", contract_item_id)
