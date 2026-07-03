"""Import reference counterparties (name, BIN/IIN, address, phone) from
parsed_counterparties.xlsx. Used to enrich debtors with contacts.

Usage (from FastAPIProject folder):
    python -m scripts.import_counterparties
    python -m scripts.import_counterparties --reset
    python -m scripts.import_counterparties --file path.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402

from app.core.database import Base, SessionLocal, engine  # noqa: E402
from app.models.counterparty import Counterparty  # noqa: E402

REPO_ROOT = BACKEND_DIR.parent
DEFAULT_XLSX = REPO_ROOT / "parsed_counterparties.xlsx"
SHEET_NAME = "counterparties"


def _clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_bin(value):
    if value is None:
        return None
    s = str(value).split(".")[0].strip()
    if not s:
        return None
    if s.isdigit() and len(s) < 12:
        return s.zfill(12)
    return s


def import_counterparties(xlsx_path: Path, reset: bool) -> int:
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        existing = db.query(Counterparty).count()
        if existing and not reset:
            print(
                f"Table 'counterparties' already has {existing} rows. "
                f"Use --reset to reimport. Aborting."
            )
            return 0
        if existing and reset:
            print(f"--reset: deleting {existing} existing rows...")
            db.query(Counterparty).delete()
            db.commit()

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb[SHEET_NAME]

        seen: set[str] = set()
        buffer: list[Counterparty] = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r is None or all(c is None for c in r):
                continue
            name, bin_iin, address, phone = (list(r) + [None] * 4)[:4]
            nb = _normalize_bin(bin_iin)
            # BIN is unique in the table; keep the first occurrence.
            if nb and nb in seen:
                continue
            if nb:
                seen.add(nb)
            buffer.append(
                Counterparty(
                    name=_clean(name),
                    bin_iin=nb,
                    address=_clean(address),
                    phone=_clean(phone),
                )
            )

        db.bulk_save_objects(buffer)
        db.commit()
        print(f"Imported {len(buffer)} counterparties.")
        return len(buffer)
    finally:
        db.close()


def main():
    parser = argparse.ArgumentParser(description="Import counterparties from xlsx.")
    parser.add_argument("--file", type=str, default=str(DEFAULT_XLSX))
    parser.add_argument("--reset", action="store_true")
    args = parser.parse_args()

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        raise SystemExit(f"File not found: {xlsx_path}")
    import_counterparties(xlsx_path, reset=args.reset)


if __name__ == "__main__":
    main()
