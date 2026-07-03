"""Import accounts-receivable (debts) from parsed_debts.xlsx into the database.

Usage (from FastAPIProject folder):
    python -m scripts.import_debts               # import if table is empty
    python -m scripts.import_debts --reset       # wipe table and reimport
    python -m scripts.import_debts --file path.xlsx

Requires DATABASE_URL in .env and openpyxl.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402

from app.core.database import Base, SessionLocal, engine  # noqa: E402
from app.models.debt import Debt  # noqa: E402

REPO_ROOT = BACKEND_DIR.parent
DEFAULT_XLSX = REPO_ROOT / "parsed_debts.xlsx"
SHEET_NAME = "debts"


def _to_float(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        try:
            return float(str(value).replace(" ", "").replace(",", "."))
        except (TypeError, ValueError):
            return None


def _clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_bin(value):
    if value is None:
        return None
    s = str(value).split(".")[0].strip()
    if not s:
        return None
    if s.isdigit() and len(s) < 12:
        return s.zfill(12)
    return s


def import_debts(xlsx_path: Path, reset: bool) -> int:
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        existing = db.query(Debt).count()
        if existing and not reset:
            print(f"Table 'debts' already has {existing} rows. Use --reset to reimport. Aborting.")
            return 0
        if existing and reset:
            print(f"--reset: deleting {existing} existing rows...")
            db.query(Debt).delete()
            db.commit()

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb[SHEET_NAME]

        buffer: list[Debt] = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r is None or all(c is None for c in r):
                continue
            (bin_iin, company_name, contract_name, turnover_debit,
             turnover_credit, debt_amount, report_date) = (list(r) + [None] * 7)[:7]

            buffer.append(
                Debt(
                    bin_iin=_normalize_bin(bin_iin),
                    company_name=_clean(company_name),
                    contract_name=_clean(contract_name),
                    turnover_debit=_to_float(turnover_debit),
                    turnover_credit=_to_float(turnover_credit),
                    debt_amount=_to_float(debt_amount),
                    report_date=_clean(report_date),
                )
            )

        db.bulk_save_objects(buffer)
        db.commit()
        print(f"Imported {len(buffer)} debt rows.")
        return len(buffer)
    finally:
        db.close()


def main():
    parser = argparse.ArgumentParser(description="Import debts from xlsx.")
    parser.add_argument("--file", type=str, default=str(DEFAULT_XLSX))
    parser.add_argument("--reset", action="store_true")
    args = parser.parse_args()

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        raise SystemExit(f"File not found: {xlsx_path}")
    import_debts(xlsx_path, reset=args.reset)


if __name__ == "__main__":
    main()
